from __future__ import annotations

import io
import json
import zipfile
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

import pandas as pd
import xlsxwriter
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

WFP_BLUE = "#007DBC"
WFP_DARK = "#005B8E"
LIGHT = "#EAF5FB"
TEXT = "#1F2937"


def _xlsx_bytes(df: pd.DataFrame, sheet_name: str, title: str, percent_cols: set[str] | None = None, status_cols: set[str] | None = None) -> bytes:
    bio = io.BytesIO()
    with xlsxwriter.Workbook(bio, {"in_memory": True}) as wb:
        ws = wb.add_worksheet(sheet_name[:31])
        title_fmt = wb.add_format({"bold": True, "font_size": 18, "font_color": "#005B8E"})
        hdr_fmt = wb.add_format({"bold": True, "font_color": "#FFFFFF", "bg_color": "#007DBC", "border": 1, "text_wrap": True, "valign": "vcenter"})
        body_fmt = wb.add_format({"border": 1, "border_color": "#D7E1E8", "valign": "top", "text_wrap": True})
        pct_fmt = wb.add_format({"border": 1, "border_color": "#D7E1E8", "num_format": "0.0%", "valign": "top"})
        date_fmt = wb.add_format({"border": 1, "border_color": "#D7E1E8", "num_format": "yyyy-mm-dd", "valign": "top"})
        kpi_label = wb.add_format({"bold": True, "font_color": "#005B8E", "bg_color": "#EAF5FB", "align": "center"})
        kpi_value = wb.add_format({"bold": True, "font_size": 14, "align": "center"})
        ws.write(0, 0, title, title_fmt)
        ws.write(2, 0, "Records", kpi_label); ws.write(3, 0, len(df), kpi_value)
        if "Severity" in df.columns:
            ws.write(2, 1, "High/Critical", kpi_label)
            n = int(df["Severity"].astype(str).str.contains("High|Critical", case=False, regex=True).sum())
            ws.write(3, 1, n, kpi_value)
        if "Status" in df.columns:
            ws.write(2, 2, "Open / Proposed", kpi_label)
            n = int(df["Status"].astype(str).isin(["Open", "Proposed"]).sum())
            ws.write(3, 2, n, kpi_value)
        start = 5
        for j, col in enumerate(df.columns):
            ws.write(start, j, col, hdr_fmt)
        percent_cols = percent_cols or set()
        for i, row in enumerate(df.itertuples(index=False, name=None), start=start+1):
            for j, value in enumerate(row):
                col = str(df.columns[j])
                fmt = pct_fmt if col in percent_cols else body_fmt
                if pd.isna(value): value = ""
                ws.write(i, j, value, fmt)
        ws.freeze_panes(start+1, 0)
        ws.autofilter(start, 0, start + max(len(df), 1), max(len(df.columns)-1, 0))
        for j, col in enumerate(df.columns):
            if col in {"Finding", "Proposed_Action", "Proposed_Management_Action", "Evidence_Summary", "Management_Comments", "Action_Basis"}:
                width = 38
            elif col in {"Detailed_Action_IDs", "Closure_Evidence"}:
                width = 30
            elif col in {"Site", "Woreda", "Activity", "Theme", "Suggested_Responsible_Unit"}:
                width = 22
            else:
                width = min(max(len(str(col)) + 2, 11), 20)
            ws.set_column(j, j, width)
        if "Severity" in df.columns and len(df):
            col = df.columns.get_loc("Severity")
            ws.conditional_format(start+1, col, start+len(df), col, {"type": "text", "criteria": "containing", "value": "Critical", "format": wb.add_format({"bg_color": "#F8D7DA", "font_color": "#842029"})})
            ws.conditional_format(start+1, col, start+len(df), col, {"type": "text", "criteria": "containing", "value": "High", "format": wb.add_format({"bg_color": "#FFF3CD", "font_color": "#664D03"})})
        if "Priority" in df.columns and len(df):
            col = df.columns.get_loc("Priority")
            ws.conditional_format(start+1, col, start+len(df), col, {"type": "text", "criteria": "containing", "value": "Critical", "format": wb.add_format({"bg_color": "#F8D7DA", "font_color": "#842029"})})
        # Lists for validation/editing columns.
        for c, values in {
            "Status": ["Proposed", "Open", "In Progress", "Pending Verification", "Closed", "On Hold"],
            "Validation_Status": ["Pending field validation", "Validated", "Corrected", "Not supported"],
            "Management_Agreement": ["Not reviewed", "Agreed", "Revise", "Not accepted"],
            "Verification_Status": ["Not started", "Pending", "Verified", "Not verified"],
        }.items():
            if c in df.columns and len(df):
                j = df.columns.get_loc(c)
                ws.data_validation(start+1, j, start+len(df), j, {"validate": "list", "source": values})
        if "Issue_Rate" in df.columns and len(df):
            j = df.columns.get_loc("Issue_Rate"); ws.set_column(j, j, 12, pct_fmt)
        if "Overall_Issue_Rate" in df.columns and len(df):
            j = df.columns.get_loc("Overall_Issue_Rate"); ws.set_column(j, j, 12, pct_fmt)
    return bio.getvalue()


def detailed_tracker_xlsx(detail: list[dict]) -> bytes:
    df = pd.DataFrame(detail)
    return _xlsx_bytes(df, "Detailed Findings", "Detailed Findings & Validation Tracker", {"Issue_Rate"})


def aggregated_tracker_xlsx(agg: list[dict]) -> bytes:
    df = pd.DataFrame(agg)
    return _xlsx_bytes(df, "Aggregated Actions", "Aggregated Management Action Tracker", {"Overall_Issue_Rate"})


def _shade(cell, fill: str):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill.replace("#", "")); tcPr.append(shd)


def _setup_doc(title: str, subtitle: str) -> Document:
    d = Document(); s = d.sections[0]
    s.top_margin=Inches(.55); s.bottom_margin=Inches(.55); s.left_margin=Inches(.65); s.right_margin=Inches(.65)
    d.styles["Normal"].font.name="Arial"; d.styles["Normal"].font.size=Pt(9); d.styles["Normal"].font.color.rgb=RGBColor.from_string("1F2937")
    for sn, size, color in [("Title", 22, "005B8E"), ("Heading 1", 15, "005B8E"), ("Heading 2", 12, "007DBC")]:
        d.styles[sn].font.name="Arial"; d.styles[sn].font.size=Pt(size); d.styles[sn].font.color.rgb=RGBColor.from_string(color); d.styles[sn].font.bold=True
    p=d.add_paragraph(); p.style="Title"; p.add_run(title).bold=True
    p=d.add_paragraph(); r=p.add_run(subtitle); r.font.color.rgb=RGBColor.from_string("64748B"); r.font.size=Pt(10)
    return d


def _table(doc: Document, headers: list[str], rows: list[list], font_size: float = 7.2):
    t=doc.add_table(rows=1, cols=len(headers)); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for j,h in enumerate(headers):
        c=t.rows[0].cells[j]; c.text=str(h); _shade(c,"007DBC")
        for rr in c.paragraphs[0].runs: rr.font.bold=True; rr.font.color.rgb=RGBColor(255,255,255); rr.font.size=Pt(font_size)
    for row in rows:
        cells=t.add_row().cells
        for j,v in enumerate(row):
            cells[j].text="" if v is None else str(v)
            for p in cells[j].paragraphs:
                for rr in p.runs: rr.font.size=Pt(font_size)
    return t


def management_report_docx(reporting_month: str, dq: dict, indicators: list[dict], agg: list[dict], detail: list[dict]) -> bytes:
    label = pd.Period(reporting_month, freq="M").strftime("%B %Y")
    d = _setup_doc("Monthly Monitoring Management Report", f"Jijiga Area Office | {label} | Automated draft — field validation required")
    d.add_heading("1. Executive overview", 1)
    p=d.add_paragraph(); p.add_run("Purpose. ").bold=True; p.add_run("This automated draft converts monthly MoDa records into structured assurance signals, site-level findings and proposed management actions. Findings are not considered agreed until field validation and internal WFP review.")
    rows=[["Accepted submissions", dq.get("rows",0)], ["Configured indicators", len(indicators)], ["Site-level findings", len(detail)], ["Aggregated actions", len(agg)]]
    _table(d,["Measure","Value"],rows,8)
    d.add_heading("2. Priority management signals",1)
    rows=[]
    for r in agg:
        rows.append([r.get("Priority"), r.get("Aggregated_Action_Title"), r.get("Activities"), r.get("Affected_Sites"), r.get("Suggested_Responsible_Unit")])
    _table(d,["Priority","Aggregated action","Activity","Sites","Suggested owner"],rows)
    d.add_heading("3. Findings by activity",1)
    byact=defaultdict(list)
    for r in indicators: byact[r.get("activity_short")].append(r)
    for a in ["A1 - Relief","A2 - Nutrition","A3 - Refugee","A6 - Resilience"]:
        rr=sorted(byact.get(a,[]), key=lambda x: (x.get("issue_rate") is not None, x.get("issue_rate") or 0), reverse=True)
        if not rr: continue
        d.add_heading(a,2)
        rows=[]
        for r in rr:
            rate = f"{100*(r.get('issue_rate') or 0):.1f}%" if r.get("applicable") else "–"
            rows.append([r.get("theme"),r.get("finding"),f"{r.get('issues',0)}/{r.get('applicable',0)}",rate,r.get("severity")])
        _table(d,["Theme","Finding","n/N","Rate","Severity"],rows)
    d.add_heading("4. Action-management workflow",1)
    for x in ["System-generated signal", "Field-team validation and context", "Internal WFP review/agreement by activity", "Owner and due date assigned", "Implementation tracked", "M&E verifies closure evidence and recurrence"]:
        d.add_paragraph(x, style="List Number")
    d.add_heading("5. Data-quality caveats",1)
    rows=[[x.get("severity"),x.get("issue"),x.get("count"),x.get("recommendation")] for x in dq.get("issues",[])]
    _table(d,["Severity","DQ issue","Count","Treatment"],rows)
    d.add_heading("6. Management follow-up",1)
    for x in ["Confirm findings after field validation.", "Assign action owners along activities/processes.", "Track open, overdue, pending-verification and closed actions monthly.", "Retain the detailed tracker as the audit trail behind aggregated management actions."]:
        d.add_paragraph(x, style="List Bullet")
    bio=io.BytesIO(); d.save(bio); return bio.getvalue()


def dq_report_docx(reporting_month: str, dq: dict, indicators: list[dict]) -> bytes:
    label = pd.Period(reporting_month, freq="M").strftime("%B %Y")
    d=_setup_doc("Monthly Monitoring Data Quality Report", f"Jijiga Area Office | {label} | Intake, schema, matching and logic checks")
    d.add_heading("1. Intake assurance",1)
    _table(d,["Measure","Value"],[["Accepted submissions",dq.get("rows",0)],["Unique UUIDs",dq.get("uuid_unique",0)],["Duplicate UUID rows removed",dq.get("duplicate_uuid_rows",0)],["Manual month mismatches",dq.get("manual_month_mismatch",0)]],8)
    d.add_heading("2. Source schema controls",1)
    rows=[]
    for sc in dq.get("schema",[]): rows.append([sc.get("file"),sc.get("header_count"),len(sc.get("matched_headers",[])),len(sc.get("missing_headers",[]))])
    _table(d,["File","Headers","Mapped configured fields","Missing configured fields"],rows)
    d.add_heading("3. Exceptions requiring review",1)
    rows=[[x.get("category"),x.get("severity"),x.get("issue"),x.get("count"),x.get("recommendation")] for x in dq.get("issues",[])]
    _table(d,["Category","Severity","Issue","Count","Recommended treatment"],rows)
    d.add_heading("4. Publication controls",1)
    for x in ["Reporting month is derived from monitoring date, not the manually selected month.", "Overlapping exports are deduplicated by _uuid.", "Questions are mapped by controlled text aliases, not Excel column position.", "Small denominators are presented as targeted assurance signals rather than generalized estimates.", "Sensitive integrity/protection signals remain verification items until independently substantiated."]:
        d.add_paragraph(x,style="List Bullet")
    bio=io.BytesIO(); d.save(bio); return bio.getvalue()


def bundle_zip(files: dict[str, bytes]) -> bytes:
    bio=io.BytesIO()
    with zipfile.ZipFile(bio,"w",zipfile.ZIP_DEFLATED) as z:
        for name, content in files.items(): z.writestr(name, content)
    return bio.getvalue()
